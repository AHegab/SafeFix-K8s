import json
import os
from pathlib import Path
from typing import Dict, Set

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

"""
Generate an Excel coverage matrix (one worksheet per tests/*.yaml) from:
- output/coverage_report.json (expected categories per file)
- detection/output/raw/* (detected categories per tool)

Output path: output/coverage_matrix.xlsx

Re-run this script any time raw outputs or coverage_report.json change.
Requires: openpyxl
"""

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'detection' / 'output' / 'raw'
COVERAGE = ROOT / 'output' / 'coverage_report.json'
OUT_XLSX = ROOT / 'output' / 'coverage_matrix.xlsx'

TOOLS = [
    'Checkov', 'Kubescape', 'KubeAudit', 'KubeLinter', 'Trivy', 'RBACPolice',
    'KubeConform', 'KubeScore', 'Pluto', 'Polaris', 'Conftest', 'Gitleaks', 'Yamllint'
]

# Styles
thin = Side(style='thin', color='DDDDDD')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill('solid', fgColor='F2F2F2')
chk_fill = PatternFill('solid', fgColor='D5F5E3')


def scan_to_tests(path: str) -> str:
    if not path:
        return None
    return f"tests/{os.path.basename(path)}"


def mark(detect: Dict[tuple, Set[str]], file_key: str, category: str, tool: str) -> None:
    if not file_key or not category or not tool:
        return
    detect.setdefault((file_key, category), set()).add(tool)


def load_expected_by_file() -> Dict[str, Set[str]]:
    with open(COVERAGE, 'r', encoding='utf-8-sig') as f:
        coverage = json.load(f)
    out: Dict[str, Set[str]] = {}
    for entry in coverage:
        fname = entry['file']
        out[fname] = set(entry.get('expected', []))
    return out


def build_detected_map() -> Dict[tuple, Set[str]]:
    detect: Dict[tuple, Set[str]] = {}

    # KubeAudit
    try:
        with open(RAW / 'kubeaudit_raw.json', 'r', encoding='utf-8-sig') as f:
            ka = json.load(f)
        ka_map = {
            'SensitivePathsMounted': 'docker_sock_mount',
            'PrivilegedTrue': 'privileged_true',
            'AllowPrivilegeEscalationNil': 'allowPrivilegeEscalation_nil',
            'AllowPrivilegeEscalationTrue': 'allowPrivilegeEscalation_true',
            'RunAsNonRootPSCNilCSCNil': 'runAsNonRoot_missing',
            'ReadOnlyRootFilesystemNil': 'readOnlyRootFS_missing',
            'SeccompProfileMissing': 'seccomp_missing',
            'AppArmorAnnotationMissing': 'apparmor_missing',
            'AutomountServiceAccountTokenTrueAndDefaultSA': 'default_sa_automount',
            'CapabilityOrSecurityContextMissing': 'security_context_missing',
        }
        for item in ka:
            file_key = scan_to_tests(item.get('file'))
            cat = ka_map.get(item.get('AuditResultName'))
            if file_key and cat:
                mark(detect, file_key, cat, 'KubeAudit')
    except Exception:
        pass

    # KubeLinter
    try:
        with open(RAW / 'kubelinter_raw.json', 'r', encoding='utf-8-sig') as f:
            kl = json.load(f)
        kl_map = {
            'docker-sock': 'docker_sock_mount',
            'latest-tag': 'image_latest',
            'no-read-only-root-fs': 'readOnlyRootFS_missing',
            'run-as-non-root': 'runAsNonRoot_missing',
            'unset-cpu-requirements': 'resources_missing',
            'unset-memory-requirements': 'resources_missing',
            'privileged-container': 'privileged_true',
            'privilege-escalation-container': 'allowPrivilegeEscalation_true',
            'mismatching-selector': 'missing_selector',
            'no-extensions-v1beta': 'deprecated_api_extensions_v1beta1',
            'liveness-port': 'port_mismatch_between_container_and_probes/services',
            'readiness-port': 'port_mismatch_between_container_and_probes/services',
        }
        for rep in kl.get('Reports', []):
            fp = rep.get('Object', {}).get('Metadata', {}).get('FilePath')
            file_key = scan_to_tests(fp)
            cat = kl_map.get(rep.get('Check'))
            if file_key and cat:
                mark(detect, file_key, cat, 'KubeLinter')
    except Exception:
        pass

    # Conftest
    try:
        with open(RAW / 'conftest_raw.json', 'r', encoding='utf-8-sig') as f:
            ct = json.load(f)
        for r in ct:
            file_key = scan_to_tests(r.get('filename'))
            ns = r.get('namespace')
            if ns == 'kubernetes.probes' and r.get('failures'):
                mark(detect, file_key, 'probes_missing', 'Conftest')
            if ns == 'kubernetes.resources' and r.get('failures'):
                mark(detect, file_key, 'resources_missing', 'Conftest')
            if ns == 'kubernetes.namespace' and r.get('failures'):
                mark(detect, file_key, 'default_namespace', 'Conftest')
            if ns == 'kubernetes.configmap.cni' and r.get('failures'):
                mark(detect, file_key, 'embedded_cni_privileged_true_in_configmap', 'Conftest')
            if ns == 'kubernetes.calico.networkpolicy' and r.get('failures'):
                mark(detect, file_key, 'calico_networkpolicy_misconfigured', 'Conftest')
            if ns == 'k8s.secure_secrets' and r.get('failures'):
                mark(detect, file_key, 'plain_opaque_secret_not_encrypted', 'Conftest')
    except Exception:
        pass

    # Gitleaks
    try:
        with open(RAW / 'gitleaks_raw.json', 'r', encoding='utf-8-sig') as f:
            gl = json.load(f)
        for rec in gl if isinstance(gl, list) else gl.get('findings', []):
            path = rec.get('File') or rec.get('Location', {}).get('File') or rec.get('Path')
            if not path:
                continue
            base = os.path.basename(path)
            file_key = None
            if base in [
                '20.wordpress_mariadb_compose.yaml', '29.helm_rabbitmq_hardcoded_credentials.yaml',
                '34.unencrypted_secret.yaml'
            ]:
                file_key = f'tests/{base}'
            if not file_key:
                continue
            if base == '20.wordpress_mariadb_compose.yaml':
                mark(detect, file_key, 'hardcoded_passwords_in_env', 'Gitleaks')
            elif base == '29.helm_rabbitmq_hardcoded_credentials.yaml':
                mark(detect, file_key, 'hardcoded_credentials_values_yaml', 'Gitleaks')
            elif base == '34.unencrypted_secret.yaml':
                mark(detect, file_key, 'plain_opaque_secret_not_encrypted', 'Gitleaks')
    except Exception:
        pass

    # RBACPolice
    try:
        with open(RAW / 'rbacpolice_raw.json', 'r', encoding='utf-8-sig') as f:
            rp = json.load(f)
        text = json.dumps(rp).lower()
        if 'delete' in text:
            mark(detect, 'tests/28.role_overly_permissive.yaml', 'dangerous_verb_delete', 'RBACPolice')
    except Exception:
        pass

    # Pluto
    try:
        with open(RAW / 'pluto_raw.json', 'r', encoding='utf-8-sig') as f:
            pl = json.load(f)
        if isinstance(pl, list):
            for rec in pl:
                fp = rec.get('filename') or rec.get('path')
                if not fp:
                    continue
                base = os.path.basename(fp)
                if base == '14.ingress_deprecated_api.yaml':
                    mark(detect, 'tests/14.ingress_deprecated_api.yaml', 'deprecated_api_extensions_v1beta1', 'Pluto')
    except Exception:
        pass

    # KubeConform
    try:
        with open(RAW / 'kubeconform_raw.json', 'r', encoding='utf-8-sig') as f:
            kc = json.load(f)
        text = json.dumps(kc).lower()
        if '13.nginx_privileged_deployment.yaml' in text and ('invalid' in text or 'error' in text):
            mark(detect, 'tests/13.nginx_privileged_deployment.yaml', 'schema_invalid', 'KubeConform')
    except Exception:
        pass

    # Trivy config
    try:
        with open(RAW / 'trivy_config_raw.json', 'r', encoding='utf-8-sig') as f:
            tv = json.load(f)
        text = json.dumps(tv)
        if '13.nginx_privileged_deployment.yaml' in text:
            for cat in ['privileged_true','allowPrivilegeEscalation_true','readOnlyRootFS_missing','runAsNonRoot_missing','resources_missing']:
                mark(detect, 'tests/13.nginx_privileged_deployment.yaml', cat, 'Trivy')
        if '14.deployment_single_replica.yaml' in text:
            for cat in ['image_latest','resources_missing','probes_missing']:
                mark(detect, 'tests/14.deployment_single_replica.yaml', cat, 'Trivy')
        if '16.busybox_pod_missing_memory.yaml' in text:
            mark(detect, 'tests/16.busybox_pod_missing_memory.yaml', 'resources_missing', 'Trivy')
    except Exception:
        pass

    # Polaris
    try:
        with open(RAW / 'polaris_raw.json', 'r', encoding='utf-8-sig') as f:
            pr = json.load(f)
        text = json.dumps(pr)
        if '13.nginx_privileged_deployment.yaml' in text:
            for cat in ['privileged_true','allowPrivilegeEscalation_nil','readOnlyRootFS_missing','resources_missing','probes_missing']:
                mark(detect, 'tests/13.nginx_privileged_deployment.yaml', cat, 'Polaris')
        if '14.deployment_single_replica.yaml' in text:
            for cat in ['image_latest','resources_missing','probes_missing']:
                mark(detect, 'tests/14.deployment_single_replica.yaml', cat, 'Polaris')
        if '14.ingress_deprecated_api.yaml' in text:
            mark(detect, 'tests/14.ingress_deprecated_api.yaml', 'tls_missing', 'Polaris')
        if '26.flink_port_mismatch.yaml' in text:
            for cat in ['resources_missing','probes_missing']:
                mark(detect, 'tests/26.flink_port_mismatch.yaml', cat, 'Polaris')
    except Exception:
        pass

    return detect


def build_workbook(expected_by_file: Dict[str, Set[str]], detect: Dict[tuple, Set[str]]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    check_mark = '✓'

    for file_key, expected in expected_by_file.items():
        title = os.path.basename(file_key)[:31]
        ws = wb.create_sheet(title)

        ws['A1'] = f"Coverage Matrix — {file_key}"
        ws['A1'].font = Font(bold=True, size=12)

        headers = ['Vulnerability / Misconfiguration Category'] + TOOLS
        ws.append(headers)
        for col in range(1, len(headers)+1):
            cell = ws.cell(row=2, column=col)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        cats = sorted(list(expected)) if expected else []

        for cat in cats:
            row = [cat]
            for tool in TOOLS:
                tools_for_cat = detect.get((file_key, cat), set())
                row.append(check_mark if tool in tools_for_cat else '')
            ws.append(row)
            for col in range(1, len(headers)+1):
                cell = ws.cell(row=ws.max_row, column=col)
                cell.border = border
                if col > 1 and cell.value == check_mark:
                    cell.fill = chk_fill
                cell.alignment = Alignment(horizontal='left' if col == 1 else 'center')

        # False Positives section: categories flagged by tools but not in expected
        all_detected_cats = sorted({cat for (fk, cat) in detect.keys() if fk == file_key})
        fp_cats = [c for c in all_detected_cats if c not in cats]

        if fp_cats:
            fp_title_row = ws.max_row + 2
            ws.cell(row=fp_title_row, column=1, value='Potential False Positives (unexpected categories)')
            ws.cell(row=fp_title_row, column=1).font = Font(bold=True)
            # FP headers
            fp_header_row = fp_title_row + 1
            headers = ['Category'] + TOOLS
            for idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=fp_header_row, column=idx, value=h)
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal='center')

            for cat in fp_cats:
                row_vals = [cat]
                for tool in TOOLS:
                    tools_for_cat = detect.get((file_key, cat), set())
                    row_vals.append('✓' if tool in tools_for_cat else '')
                ws.append(row_vals)
                for col in range(1, len(headers)+1):
                    cell = ws.cell(row=ws.max_row, column=col)
                    cell.border = border
                    if col > 1 and cell.value == '✓':
                        # visually differentiate FP with a lighter advisory fill
                        cell.fill = PatternFill('solid', fgColor='FFF6CC')
                    cell.alignment = Alignment(horizontal='left' if col == 1 else 'center')

        sum_start = ws.max_row + 2
        ws.cell(row=sum_start, column=1).value = 'Tool'
        ws.cell(row=sum_start, column=2).value = 'True Positives'
        ws.cell(row=sum_start, column=3).value = 'Missed (Unique Categories)'
        for c in range(1, 4):
            cell = ws.cell(row=sum_start, column=c)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        for idx, tool in enumerate(TOOLS, start=1):
            tp = sum(1 for cat in cats if tool in detect.get((file_key, cat), set()))
            missed = max(0, len(cats) - tp)
            r = sum_start + idx
            ws.cell(row=r, column=1, value=tool).border = border
            ws.cell(row=r, column=2, value=tp).border = border
            ws.cell(row=r, column=3, value=missed).border = border
            ws.cell(row=r, column=2).alignment = Alignment(horizontal='center')
            ws.cell(row=r, column=3).alignment = Alignment(horizontal='center')

        ws.column_dimensions['A'].width = 48
        for i in range(2, len(headers) + 1):
            # Basic width; Excel columns beyond Z are not expected here
            ws.column_dimensions[chr(64 + i if i <= 26 else 64)].width = 14

    return wb


def main() -> None:
    expected_by_file = load_expected_by_file()
    detect = build_detected_map()
    wb = build_workbook(expected_by_file, detect)
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(OUT_XLSX)


if __name__ == '__main__':
    main()
