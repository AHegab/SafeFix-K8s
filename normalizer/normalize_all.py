# normalizer/normalize_all.py
# Robust normalizer for SafeFixK8s outputs
from __future__ import annotations
import json, hashlib, sys, re
from pathlib import Path
from datetime import datetime
from typing import Any, Dict, Iterable, List

ROOT    = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "detection" / "output" / "raw"
OUT     = ROOT / "output" / "normalized_findings.json"

# ------------------------
# helpers
# ------------------------

def now_iso() -> str:
    return datetime.utcnow().isoformat(timespec="microseconds") + "Z"

def warn(msg: str) -> None:
    print(f"[WARN] {msg}")

def _hash_id(parts: Iterable[str]) -> str:
    h = hashlib.md5()
    for p in parts:
        if p is None:
            p = ""
        if not isinstance(p, str):
            p = json.dumps(p, sort_keys=True, ensure_ascii=False)
        h.update(p.encode("utf-8", errors="ignore"))
        h.update(b"|")
    return h.hexdigest()

def from_kubeaudit(p: Path) -> List[Dict]:
    data = load_json(p)
    findings: List[Dict] = []
    items = data if isinstance(data, list) else ([data] if isinstance(data, dict) else [])
    for it in items:
        if not isinstance(it, dict):
            continue
        rule = it.get("auditResultName") or it.get("name") or it.get("id") or "kubeaudit"
        sev = (it.get("severity") or "warning").upper()
        if sev == "ERROR": sev = "HIGH"
        elif sev == "WARNING": sev = "MEDIUM"
        elif sev == "INFO": sev = "LOW"
        obj = it.get("resource") or {}
        kind = obj.get("kind") or "Object"
        name = obj.get("name") or "unknown"
        ns = obj.get("namespace")
        msg = it.get("message") or it.get("remediation") or rule
        labels = {}
        if ns: labels["namespace"] = ns
        findings.append(_mk("kubeaudit", rule, sev, kind, name, msg, labels=labels))
    return findings

def load_json(p: Path) -> Any:
    """
    Robust loader:
    - UTF-8 with optional BOM
    - Accepts JSON arrays/objects
    - If JSONL, returns a list of decoded objects
    - If file is placeholder / not valid JSON, returns []
    """
    try:
        if not p.exists() or p.stat().st_size == 0:
            return []
        raw = p.read_bytes()
        # Strip UTF-8 BOM if present
        if raw.startswith(b"\xef\xbb\xbf"):
            raw = raw[3:]

        text = raw.decode("utf-8", errors="replace").strip()
        if not text:
            return []

        # Skip obvious placeholders or shell error text
        low = text.lower()
        if ("placeholder" in low) or ("error response from daemon" in low) or ("exec /kubeaudit" in low):
            warn(f"Placeholder/diagnostic text in {p}; skipping.")
            return []

        # First try standard JSON
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            # Try JSON Lines
            objs: List[Any] = []
            ok = True
            for line in text.splitlines():
                line = line.strip()
                if not line:
                    continue
                try:
                    objs.append(json.loads(line))
                except json.JSONDecodeError:
                    ok = False
                    break
            if ok and objs:
                return objs

            warn(f"Bad JSON in {p}: falling back to empty.")
            return []
    except Exception as e:
        warn(f"Failed to read {p}: {e}")
        return []
    
def iter_json_lines(p: Path) -> Iterable[Any]:
    """Yield JSON objects from a file that may contain JSON-lines with noise."""
    if not p.exists() or p.stat().st_size == 0:
        warn(f"Missing/empty: {p}")
        return
    with p.open("r", encoding="utf-8-sig", errors="ignore") as f:
        for ln in f:
            s = ln.strip()
            if not s:
                continue
            # ignore Go log lines like {"level":"info-0",...} that *are* JSON too – still parse them
            try:
                yield json.loads(s)
            except Exception:
                # skip non-JSON garbage lines
                continue

def safe_text(p: Path) -> str:
    if not p.exists() or p.stat().st_size == 0:
        warn(f"Missing/empty: {p}")
        return ""
    try:
        return p.read_text(encoding="utf-8-sig", errors="ignore")
    except Exception as e:
        warn(f"Text read failed {p}: {e}")
        return ""
    
def sev_norm(s: Any, default="INFO") -> str:
    if s is None:
        return default
    s = str(s)
    su = s.upper()
    sl = s.lower()
    if su in {"LOW","MEDIUM","HIGH","CRITICAL","INFO","UNKNOWN"}:
        return su
    if sl == "critical": return "CRITICAL"
    if sl in {"warning","warn"}: return "MEDIUM"
    if sl in {"advice","info","ok","passed","bestpractice","best_practice"}: return "INFO"
    if su.startswith("WARN"): return "MEDIUM"
    if su.startswith("ERR"): return "HIGH"
    return default


def _norm_msg(msg: Any) -> str:
    if isinstance(msg, dict):
        for k in ("Message", "message", "summary", "text"):
            v = msg.get(k)
            if isinstance(v, str):
                return v
        return json.dumps(msg, sort_keys=True, ensure_ascii=False)
    if msg is None:
        return ""
    return str(msg)

def _dedupe_key(f: Dict[str, Any]) -> tuple:
    rr = f.get("resourceRef", {}) or {}
    loc = f.get("location", {}) or {}
    return (
        f.get("tool",""),
        f.get("ruleId",""),
        rr.get("kind",""),
        rr.get("name",""),
        loc.get("file",""),
        int(loc.get("line", 0) or 0),
        _norm_msg(f.get("message"))
    )

def _mk(
    tool: str,
    rule: str | None,
    severity: str | None,
    kind: str | None,
    name: str | None,
    message: Any,
    file: str | None = None,
    line: int | None = None,
    labels: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    msg = _norm_msg(message)
    fid = _hash_id([tool, rule, kind, name, file, str(line or 0), msg])
    return {
        "findingId": fid,
        "tool": tool,
        "ruleId": rule or "unknown",
        "severity": sev_norm(severity, default="INFO"),
        "resourceRef": {"kind": kind or "Object", "name": name or "unknown"},
        "message": msg if isinstance(message, str) else message,
        "location": {"file": file or "", "line": line},
        "labels": labels or {}
    }

# ------------------------
# mappers
# ------------------------

def from_kubelinter(p: Path) -> List[Dict]:
    data = load_json(p) or {}
    findings = []
    for it in data.get("Reports", []) or data.get("reports", []) or []:
        if not isinstance(it, dict):
            continue
        # kube-linter JSON v2/v1 vary; try generic fields
        tool = "kubelinter"
        # Check can be a string directly or nested object
        check_field = it.get("Check")
        if isinstance(check_field, dict):
            rule = check_field.get("name") or check_field.get("id")
        else:
            rule = check_field or it.get("check")
        
        # Get message from Diagnostic or directly
        diag = it.get("Diagnostic", {})
        if isinstance(diag, dict):
            msg = diag.get("Message") or diag.get("message")
        else:
            msg = it.get("Message") or it.get("message")
        
        sev  = it.get("Severity") or it.get("severity") or "MEDIUM"
        obj  = it.get("Object", {}) or it.get("object", {}) or {}
        
        # Handle K8sObject nested structure
        k8s_obj = obj.get("K8sObject") or obj.get("k8sObject") or {}
        if isinstance(k8s_obj, dict):
            gvk = k8s_obj.get("GroupVersionKind") or {}
            kind = gvk.get("Kind") or k8s_obj.get("kind") or "Object"
            name = k8s_obj.get("Name") or k8s_obj.get("name") or "unknown"
        else:
            kind = obj.get("kind") or "Object"
            name = obj.get("name") or "unknown"
        
        # Get file from Metadata
        metadata = obj.get("Metadata") or obj.get("metadata") or {}
        file = metadata.get("FilePath") or metadata.get("filePath") or obj.get("file") or ""
        
        findings.append(_mk(tool, rule, sev, kind, name, {"Message": msg or rule}, file=file))
    return findings

def from_polaris(p: Path) -> List[Dict]:
    data = load_json(p) or {}
    findings = []
    # Polaris structure: Results -> each has Results (checks) and PodResult
    for res in data.get("Results", []) or data.get("results", []):
        if not isinstance(res, dict):
            continue
        kind = res.get("Kind") or res.get("kind") or "Object"
        name = res.get("Name") or res.get("name") or "unknown"
        namespace = res.get("Namespace") or res.get("namespace") or ""
        file = res.get("FilePath") or res.get("filePath") or ""
        
        # Process resource-level checks
        checks_dict = res.get("Results", {}) or res.get("results", {})
        if isinstance(checks_dict, dict):
            for check_id, check_data in checks_dict.items():
                if not isinstance(check_data, dict):
                    continue
                # Only report failures (Success: false)
                if check_data.get("Success", True):
                    continue
                rule = check_data.get("ID") or check_id
                sev = check_data.get("Severity") or check_data.get("severity") or "INFO"
                # Map polaris severity: danger->HIGH, warning->MEDIUM
                if sev.lower() == "danger":
                    sev = "HIGH"
                elif sev.lower() == "warning":
                    sev = "MEDIUM"
                msg = check_data.get("Message") or check_data.get("message") or rule
                findings.append(_mk("polaris", rule, sev, kind, name, msg, file=file))
        
        # Process pod-level checks
        pod_result = res.get("PodResult") or res.get("podResult")
        if isinstance(pod_result, dict):
            pod_checks = pod_result.get("Results", {}) or pod_result.get("results", {})
            if isinstance(pod_checks, dict):
                for check_id, check_data in pod_checks.items():
                    if not isinstance(check_data, dict):
                        continue
                    if check_data.get("Success", True):
                        continue
                    rule = check_data.get("ID") or check_id
                    sev = check_data.get("Severity") or check_data.get("severity") or "INFO"
                    if sev.lower() == "danger":
                        sev = "HIGH"
                    elif sev.lower() == "warning":
                        sev = "MEDIUM"
                    msg = check_data.get("Message") or check_data.get("message") or rule
                    findings.append(_mk("polaris", rule, sev, kind, name, msg, file=file))
    return findings

def from_trivy_config(p: Path) -> List[Dict]:
    data = load_json(p) or {}
    findings = []
    for r in data.get("Results", []) or data.get("results", []) or []:
        file = r.get("Target") or r.get("target") or ""
        for m in r.get("Misconfigurations", []) or r.get("misconfigurations", []) or []:
            rule = m.get("ID") or m.get("id")
            sev  = m.get("Severity") or m.get("severity")
            msg  = m.get("Message") or m.get("message")
            kind = "Object"
            name = "unknown"
            findings.append(_mk("trivy-config", rule, sev, kind, name, msg, file=file, labels={"url": m.get("PrimaryURL") or m.get("primaryURL")}))
    return findings

def from_kubescore(p: Path) -> List[Dict]:
    data = load_json(p)
    findings: List[Dict] = []
    if not data:
        return findings

    items = data if isinstance(data, list) else [data]

    for item in items:
        objmeta = item.get("object", {}) if isinstance(item, dict) else {}
        kind = (objmeta.get("kind") or item.get("kind") or "Object") if isinstance(item, dict) else "Object"
        name = (objmeta.get("name") or item.get("name") or "unknown") if isinstance(item, dict) else "unknown"
        file = ""
        if isinstance(item, dict):
            file = item.get("fileName") or item.get("filename") or item.get("file") or ""

        checks = []
        if isinstance(item, dict) and isinstance(item.get("checks"), list):
            checks = item["checks"]
        elif isinstance(item, dict) and any(k in item for k in ("check","id","grade","severity","comment","comments","message")):
            checks = [item]

        for chk in checks:
            rule = (chk.get("check") or chk.get("id") or chk.get("name") or "unknown")
            raw_sev = (chk.get("grade") or chk.get("severity") or chk.get("type") or "MEDIUM")
            if str(raw_sev).lower() == "security":
                raw_sev = "HIGH"
            msg = (chk.get("comment") or chk.get("comments") or chk.get("message") or chk)
            findings.append(_mk("kube-score", rule, raw_sev, kind, name, msg, file=file))

    return findings

_yamllint_re = re.compile(r"^(?P<file>.+):(?P<line>\d+):(?P<col>\d+):\s\[(?P<severity>\w+)\]\s(?P<message>.+)$")

def from_yamllint(p: Path) -> List[Dict]:
    txt = safe_text(p)
    findings = []
    for ln in txt.splitlines():
        m = _yamllint_re.match(ln.strip())
        if not m:
            # try braces warning line (tool sometimes prints different)
            continue
        d = m.groupdict()
        findings.append(_mk(
            "yamllint",
            rule="yaml-lint",
            severity=d.get("severity","LOW"),
            kind="YAML",
            name=Path(d["file"]).name,
            message=d["message"],
            file=d["file"],
            line=int(d["line"])
        ))
    return findings

def from_checkov(p: Path) -> List[Dict]:
    data = load_json(p) or {}
    findings: List[Dict] = []

    # Checkov may return an object OR a list of report objects
    items: List[Dict] = []
    if isinstance(data, list):
        items = [d for d in data if isinstance(d, dict)]
    elif isinstance(data, dict):
        items = [data]

    for item in items:
        res = (item.get("results") or {}) if isinstance(item, dict) else {}
        for r in (res.get("failed_checks") or []):
            rule = r.get("check_id")
            sev  = r.get("severity") or "INFO"
            msg  = r.get("check_name") or r.get("description")
            addr = r.get("resource") or r.get("resource_address") or ""
            kind = "Object"
            name = addr or "unknown"
            file = r.get("file_path") or r.get("repo_file_path") or ""
            line = (r.get("file_line_range") or [None, None])[0]
            labels = {"guideline": r.get("guideline")}
            findings.append(_mk("checkov", rule, sev, kind, name, msg, file=file, line=line, labels=labels))
    return findings


def from_terrascan(p: Path) -> List[Dict]:
    data = load_json(p) or {}
    findings = []
    for v in data.get("results", {}).get("violations", []):
        rule = v.get("rule_name") or v.get("rule_id") or v.get("violation_id")
        sev  = v.get("severity") or "MEDIUM"
        msg  = v.get("description") or v.get("category") or v.get("violation_message") or v
        kind = v.get("resource", {}).get("type") or "Object"
        name = v.get("resource", {}).get("name") or "unknown"
        file = v.get("file") or v.get("file_path") or ""
        line = v.get("line")
        findings.append(_mk("terrascan", rule, sev, kind, name, msg, file=file, line=line))
    return findings

def from_gitleaks(p: Path) -> List[Dict]:
    """
    Parse gitleaks JSON output.
    Gitleaks outputs a mix of human-readable text and JSON array.
    We need to extract the JSON array part.
    
    Example structure:
    [
      {
        "RuleID": "kubernetes-secret-yaml",
        "Description": "Possible Kubernetes Secret detected...",
        "StartLine": 2,
        "EndLine": 9,
        "Match": "kind: Secret\\n...",
        "Secret": "password: cGFzc3dvcmQ=",
        "File": "/scan/tests/34.unencrypted_secret.yaml",
        "Fingerprint": "/scan/tests/34.unencrypted_secret.yaml:kubernetes-secret-yaml:2"
      }
    ]
    """
    findings = []
    
    # Read file and extract JSON array
    try:
        if not p.exists() or p.stat().st_size == 0:
            return findings
        
        # Read as bytes first to handle BOM properly
        raw_bytes = p.read_bytes()
        
        # Strip UTF-8 BOM if present
        if raw_bytes.startswith(b"\xef\xbb\xbf"):
            raw_bytes = raw_bytes[3:]
        
        content = raw_bytes.decode("utf-8", errors="ignore")
        
        # Find the JSON array part - look for newline + [ + newline + { pattern
        # to avoid matching ANSI color codes like [1;3;m
        import re
        match = re.search(r'\n\[\s*\n\s*\{', content)
        if not match:
            warn(f"No JSON array found in {p.name}")
            return findings
        
        # Extract from the matched '[' position
        json_start = match.start() + 1  # +1 to skip the \n before [
        json_part = content[json_start:].strip()
        data = json.loads(json_part)
        
        if not isinstance(data, list):
            data = [data]
        
        for item in data:
            if not isinstance(item, dict):
                continue
            
            # Skip if not a finding (log lines)
            rule_id = item.get("RuleID")
            if not rule_id:
                continue
            
            description = item.get("Description", "")
            file_path = item.get("File", "")
            start_line = item.get("StartLine", 0)
            secret_match = item.get("Secret", "")
            
            # Extract filename from path
            file_name = file_path.split("/")[-1] if file_path else "unknown"
            
            # Severity based on rule type
            severity = "HIGH"
            if "generic" in rule_id.lower():
                severity = "MEDIUM"
            
            message = description or f"Secret detected: {secret_match[:50]}"
            
            findings.append(_mk(
                "gitleaks",
                rule_id,
                severity,
                "Secret",
                file_name,
                message,
                file=file_path,
                line=int(start_line)
            ))
    
    except Exception as e:
        warn(f"Failed to parse {p.name}: {e}")
    
    return findings

def from_trufflehog(p: Path) -> List[Dict]:
    # trufflehog docker emits a mix of logs (JSON) and results as JSON lines
    findings = []
    for it in iter_json_lines(p):
        if not isinstance(it, dict):
            continue
        # Known fields: "DetectorName" / "Raw", "Redacted", "SourceMetadata", "Verification"
        detector = it.get("DetectorName") or it.get("DetectorType") or it.get("rule")
        if not detector:
            # might be a log line like {"level":"info-0", ...}
            continue
        sev = "HIGH"
        msg = f"High-entropy or credential-like pattern"
        src = (((it.get("SourceMetadata") or {}).get("Data") or {}).get("Filesystem") or {})
        file = src.get("file") or it.get("File") or ""
        line = src.get("line") or it.get("Line") or 0
        rule = detector
        findings.append(_mk("trufflehog", rule, sev, "Repository", "\\", msg, file=file, line=int(line or 0)))
    return findings

def from_syft(p: Path) -> List[Dict]:
    data = load_json(p) or {}
    findings = []
    for pkg in data.get("artifacts", []) or data.get("artifactsByImage", []) or []:
        name = pkg.get("name")
        ver  = pkg.get("version")
        labels = {"type": (pkg.get("type") or pkg.get("packageType") or "").lower()}
        findings.append(_mk("syft","SBOM-PACKAGE","INFO","Image", f"{name}:{ver}", f"Package {name}:{ver} present", labels=labels))
    # some older syft formats:
    for img in data.get("artifacts", []) if isinstance(data.get("artifacts"), dict) else []:
        pass
    return findings

def from_grype(p: Path) -> List[Dict]:
    data = load_json(p) or {}
    findings = []
    for m in data.get("matches", []) or []:
        vuln = (m.get("vulnerability") or {})
        art  = (m.get("artifact") or {})
        rule = vuln.get("id") or "unknown"
        sev  = vuln.get("severity") or "MEDIUM"
        pkg  = f"{art.get('name')} {art.get('version')}"
        img  = (data.get("source", {}) or {}).get("target", "") or "image"
        msg  = f"{pkg} vulnerable"
        findings.append(_mk("grype", rule, sev, "Image", img, msg))
    return findings

def from_kubescape(p: Path) -> List[Dict]:
    data = load_json(p) or {}
    findings: List[Dict] = []

    # If it's an array of per-file results, flatten and keep only failures
    if isinstance(data, list):
        for entry in data:
            if not isinstance(entry, dict):
                continue
            # try standard fields
            results = entry.get("results") or []
            resources = entry.get("resources") or []
            # build map resourceID -> (kind,name,file)
            rmap = {}
            for res in resources:
                rid = res.get("resourceID", "")
                if not rid: continue
                obj = res.get("object", {}) or {}
                src = res.get("source", {}) or {}
                rmap[rid] = {
                    "kind": obj.get("kind",""),
                    "name": (obj.get("metadata", {}) or {}).get("name",""),
                    "file": src.get("relativePath",""),
                }
            for resu in results:
                rid = resu.get("resourceID","")
                meta = rmap.get(rid, {})
                for ctrl in resu.get("controls") or []:
                    status = (ctrl.get("status") or {}).get("status")
                    if status != "failed":
                        continue
                    cid = ctrl.get("controlID","")
                    cname = ctrl.get("name","")
                    msg = cname or cid or "Kubescape control failed"
                    findings.append(_mk(
                        "kubescape",
                        cid or "unknown",
                        "MEDIUM",
                        meta.get("kind","Resource"),
                        meta.get("name","unknown"),
                        msg,
                        file=meta.get("file",""),
                    ))
        return findings

    # otherwise, handle the single-object schema as before
    controls_summary = (data.get("summaryDetails", {}) or {}).get("controls", {}) or {}
    resources_list = data.get("resources", []) or []
    resource_map = {}
    for res in resources_list:
        res_id = res.get("resourceID", "")
        if res_id:
            obj = res.get("object", {}) or {}
            source = res.get("source", {}) or {}
            resource_map[res_id] = {
                "kind": obj.get("kind", ""),
                "name": obj.get("metadata", {}).get("name", "") if isinstance(obj.get("metadata"), dict) else "",
                "file": source.get("relativePath", "")
            }

    results = data.get("results", []) or []
    for result in results:
        resource_id = result.get("resourceID", "")
        controls = result.get("controls", []) or []
        res_meta = resource_map.get(resource_id, {})
        kind = res_meta.get("kind", "Resource")
        name = res_meta.get("name", "unknown")
        file = res_meta.get("file", "")
        for control in controls:
            status_info = control.get("status", {}) or {}
            if status_info.get("status") != "failed":
                continue
            control_id = control.get("controlID", "")
            control_name = control.get("name", "")
            sev = "MEDIUM"
            findings.append(_mk("kubescape", control_id, sev, kind, name, control_name or control_id, file=file))
    return findings

def from_kubesec(p: Path) -> List[Dict]:
    data = load_json(p)
    findings = []
    if isinstance(data, list):
        items = data
    elif isinstance(data, dict):
        items = data.get("results") or data.get("checks") or []
        if not isinstance(items, list):
            items = [items]
    else:
        items = []
    for it in items:
        if not isinstance(it, dict):
            continue
        rule = it.get("id") or it.get("selector") or "kubesec"
        sev  = it.get("severity") or "MEDIUM"
        file = (it.get("file") or it.get("target") or it.get("path") or "")
        obj = it.get("object", {})
        if isinstance(obj, dict):
            kind = obj.get("kind") or "Object"
            name = obj.get("name") or "unknown"
        else:
            kind = "Object"
            name = "unknown"
        msg  = it.get("reason") or it.get("message") or it.get("recommendations") or it
        findings.append(_mk("kubesec", rule, sev, kind, name, msg, file=file))
    return findings

# ------------------------
# collector
# ------------------------

def _fingerprint(f: Dict) -> str:
    """
    Stable hash for deduping: based on sorted JSON of a reduced view
    so that dicts are never used as set keys.
    """
    reduced = {
        "tool": f.get("tool", ""),
        "ruleId": f.get("ruleId", ""),
        "severity": f.get("severity", ""),
        "resourceKind": f.get("resourceRef", {}).get("kind", ""),
        "resourceName": f.get("resourceRef", {}).get("name", ""),
        "file": f.get("location", {}).get("file", ""),
        "line": f.get("location", {}).get("line", ""),
        "message": str(f.get("message", ""))[:200],
    }
    blob = json.dumps(reduced, sort_keys=True, ensure_ascii=False, default=str)
    return hashlib.sha1(blob.encode("utf-8")).hexdigest()

def collect_all() -> List[Dict]:
    findings: List[Dict] = []
    findings += from_kubelinter(RAW_DIR / "kubelinter_raw.json")
    findings += from_kubeaudit(RAW_DIR / "kubeaudit_raw.json")
    findings += from_polaris(RAW_DIR / "polaris_raw.json")
    findings += from_trivy_config(RAW_DIR / "trivy_config_raw.json")
    findings += from_kubescore(RAW_DIR / "kubescore_raw.json")
    findings += from_yamllint(RAW_DIR / "yamllint_raw.txt")
    findings += from_checkov(RAW_DIR / "checkov_raw.json")
    findings += from_terrascan(RAW_DIR / "terrascan_raw.json")
    findings += from_gitleaks(RAW_DIR / "gitleaks_raw.json")
    findings += from_trufflehog(RAW_DIR / "trufflehog_raw.json")
    # Optional extras; load_json returns [] for bad placeholders, so these are safe:
    findings += from_kubescape(RAW_DIR / "kubescape_raw.json")
    findings += from_kubesec(RAW_DIR / "kubesec_raw.json")

    seen: set[str] = set()
    deduped: List[Dict] = []
    for f in findings:
        key = _fingerprint(f)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(f)
    return deduped

def generate_excel_report(findings: List[Dict]) -> None:
    """Generate Excel report with tools as columns and vulnerabilities grouped by resource."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        warn("openpyxl not installed. Run: pip install openpyxl")
        warn("Skipping Excel generation...")
        return
    
    # Group findings by resource (kind + name + file)
    grouped = {}
    for f in findings:
        res_ref = f.get("resourceRef", {})
        loc = f.get("location", {})
        resource_key = (
            res_ref.get("kind", "Object"),
            res_ref.get("name", "unknown"),
            loc.get("file", "")
        )
        if resource_key not in grouped:
            grouped[resource_key] = {}
        
        tool = f.get("tool", "unknown")
        rule_raw = f.get("ruleId", "unknown")
        # Convert rule to string if it's a dict
        if isinstance(rule_raw, dict):
            rule = rule_raw.get("id") or rule_raw.get("name") or json.dumps(rule_raw)
        else:
            rule = str(rule_raw)
        
        severity = f.get("severity", "INFO")
        message = f.get("message", "")
        if isinstance(message, dict):
            message = message.get("Message", json.dumps(message))
        message = str(message)[:200]  # Limit message length
        
        vuln_key = f"{rule} [{severity}]"
        if vuln_key not in grouped[resource_key]:
            grouped[resource_key][vuln_key] = {
                "rule": rule,
                "severity": severity,
                "message": message,
                "tools": set()
            }
        grouped[resource_key][vuln_key]["tools"].add(tool)
    
    # Get all unique tools
    all_tools = set()
    for resource_vulns in grouped.values():
        for vuln_data in resource_vulns.values():
            all_tools.update(vuln_data["tools"])
    all_tools = sorted(all_tools)
    
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Security Findings"
    
    # Define styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers
    headers = ["Resource Kind", "Resource Name", "File", "Vulnerability", "Severity", "Message"] + all_tools
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    
    # Data rows
    row_idx = 2
    for (kind, name, file), vulns in sorted(grouped.items()):
        resource_start_row = row_idx
        
        for vuln_key, vuln_data in sorted(vulns.items()):
            ws.cell(row=row_idx, column=1, value=kind).border = border
            ws.cell(row=row_idx, column=2, value=name).border = border
            ws.cell(row=row_idx, column=3, value=file or "N/A").border = border
            ws.cell(row=row_idx, column=4, value=vuln_data["rule"]).border = border
            
            # Severity cell with color
            sev_cell = ws.cell(row=row_idx, column=5, value=vuln_data["severity"])
            sev_cell.border = border
            if vuln_data["severity"] == "CRITICAL":
                sev_cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
                sev_cell.font = Font(color="FFFFFF", bold=True)
            elif vuln_data["severity"] == "HIGH":
                sev_cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
                sev_cell.font = Font(color="FFFFFF")
            elif vuln_data["severity"] == "MEDIUM":
                sev_cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
            elif vuln_data["severity"] == "LOW":
                sev_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            
            ws.cell(row=row_idx, column=6, value=vuln_data["message"]).border = border
            
            # Tool columns - mark with "yes", "no", or leave empty
            for col_idx, tool in enumerate(all_tools, 7):
                cell = ws.cell(row=row_idx, column=col_idx)
                if tool in vuln_data["tools"]:
                    cell.value = "yes"
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                else:
                    cell.value = "no"
                cell.border = border
            
            row_idx += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 50
    for col_idx in range(7, 7 + len(all_tools)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 12
    
    # Freeze header row
    ws.freeze_panes = "A2"
    
    # Save
    excel_path = OUT.parent / "security_findings.xlsx"
    wb.save(excel_path)
    print(f"[ok] wrote Excel report: {excel_path}")

def main():
    all_findings = collect_all()
    
    bundle = {
        "version": "1.0",
        "generatedAt": now_iso(),
        "source": "SafeFixK8s/Normalizer",
        "findings": all_findings
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8") as f:
        json.dump(bundle, f, ensure_ascii=False, indent=2)
    print(f"[ok] wrote {OUT} with {len(bundle['findings'])} findings.")
    
    # Generate Excel report
    generate_excel_report(all_findings)
    
    # NO LONGER DELETING RAW FILES - Keep them for reference
    print(f"[info] Raw detection files preserved in: {RAW_DIR}")

if __name__ == "__main__":
    main()
